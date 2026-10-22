#!/usr/bin/env python3
"""Prepare scoped Kalodata exports without network calls or database writes.

The manifest is a JSON array of {filename, scope, source_id, source_url, period}.
Supported scopes: product_videos and shop_products. IDs remain strings. The
existing TypeScript normalizer validates commercial fields, but its row rank
is kept only as source_rank; these files are never full catalogue rankings.

Run with the bundled Python that includes openpyxl. Review migration.sql and
apply.sql independently before execution. Neither is executed by this tool.
"""
from __future__ import annotations

import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re
import subprocess
from urllib.parse import parse_qs, urlparse
import warnings

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SCOPES = {"product_videos": ("videos", "LIST_VIDEO", "product"),
          "shop_products": ("products", "LIST_PRODUCT", "shop")}
ID = re.compile(r"^\d{15,22}$")


def dump(path: Path, value) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")


def period_dates(value: str) -> tuple[date, date]:
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}~\d{4}-\d{2}-\d{2}", value):
        raise ValueError("Expected an exact YYYY-MM-DD~YYYY-MM-DD report period")
    start, end = map(date.fromisoformat, value.split("~"))
    if end < start:
        raise ValueError("Report end precedes its start")
    return start, end


def source_identity(url: str, entity: str) -> str:
    parsed = urlparse(url)
    if (parsed.scheme != "https" or parsed.hostname not in ("www.kalodata.com", "kalodata.com")
            or parsed.username or parsed.password or parsed.port or parsed.fragment
            or parsed.path.rstrip("/") != f"/{entity}/detail"):
        raise ValueError(f"Expected an observed Kalodata {entity} detail URL")
    ids = parse_qs(parsed.query).get("id", [])
    if len(ids) != 1 or not ID.fullmatch(ids[0]):
        raise ValueError("Source URL has no unambiguous numeric ID")
    return ids[0]


def clean_cell(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def read_workbook(path: Path, sheet_name: str):
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError(f"{path.name}: missing {sheet_name}")
        rows = list(book[sheet_name].values)
        headers = rows[0] if rows else ()
        if not headers or any(not isinstance(h, str) or not h for h in headers) or len(headers) != len(set(headers)):
            raise ValueError(f"{path.name}: blank or duplicate column headers")
        source_rows = [dict(zip(headers, map(clean_cell, row))) for row in rows[1:] if any(v is not None for v in row)]
        intro = [[clean_cell(v) for v in row] for row in book['Intro'].values] if 'Intro' in book.sheetnames else []
        return source_rows, intro
    finally:
        book.close()


def normalize(rows, kind, market, node):
    module = (ROOT / 'supabase/functions/_shared/kalodata-import.ts').as_uri()
    script = f'''import {{normalizeKalodata}} from {json.dumps(module)};
let text=''; for await (const part of process.stdin) text+=part;
const input=JSON.parse(text);
process.stdout.write(JSON.stringify(normalizeKalodata(input.rows,input.kind,input.market)));'''
    result = subprocess.run([node, '--experimental-strip-types', '--input-type=module', '-e', script],
                            input=json.dumps(dict(rows=rows, kind=kind, market=market)),
                            text=True, capture_output=True, timeout=60)
    if result.returncode:
        raise ValueError(f"Existing Kalodata normalizer rejected input: {result.stderr.strip()}")
    return json.loads(result.stdout)


def prepare_export(entry, base: Path, market: str, node: str):
    required = {'filename', 'scope', 'source_id', 'source_url', 'period'}
    if not isinstance(entry, dict) or not required <= entry.keys():
        raise ValueError(f"Every manifest entry needs {sorted(required)}")
    if entry['scope'] not in SCOPES or not isinstance(entry['source_id'], str) or not ID.fullmatch(entry['source_id']):
        raise ValueError("Invalid scope or source_id; numeric IDs must be strings")
    kind, sheet, entity = SCOPES[entry['scope']]
    if source_identity(entry['source_url'], entity) != entry['source_id']:
        raise ValueError("Manifest source ID and source URL disagree")
    period_dates(entry['period'])
    path = Path(entry['filename']).expanduser()
    if not path.is_absolute():
        path = base / path
    path = path.resolve(strict=True)
    raw, intro = read_workbook(path, sheet)
    normalized = normalize(raw, kind, market, node)
    if normalized['period'] != entry['period']:
        raise ValueError(f"{path.name}: manifest period differs from Excel")
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    rows = []
    for index, (record, source) in enumerate(zip(normalized['rows'], raw), start=1):
        if kind == 'videos' and record.get('tiktok_product_id') != entry['source_id']:
            raise ValueError(f"{path.name}, row {index + 1}: video has a different source product")
        key = record['tiktok_video_id' if kind == 'videos' else 'tiktok_product_id']
        if not isinstance(key, str) or not ID.fullmatch(key):
            raise ValueError(f"{path.name}: invalid item ID")
        # A product/creator page rank must never become a global catalogue rank.
        source_rank = record.pop('rank', index)
        rows.append(dict(source_record_id=key, source_rank=source_rank,
                         normalized=record, source_row=source))
    return dict(filename=str(path), source_filename=path.name, source_sha256=sha,
                scope=entry['scope'], source_id=entry['source_id'], source_url=entry['source_url'],
                market=market, period=normalized['period'], intro=intro, rows=rows)


MIGRATION = '''-- Review/apply as a migration before apply.sql. No public raw data.
create table public.kalodata_targeted_report_rows (
  market text not null check (market in ('mx','us')),
  source_scope text not null check (source_scope in ('product_videos','shop_products')),
  source_entity_id text not null check (source_entity_id ~ '^[0-9]{15,22}$'),
  report_date_range text not null check (report_date_range ~ '^\\d{4}-\\d{2}-\\d{2}~\\d{4}-\\d{2}-\\d{2}$'),
  source_record_id text not null check (source_record_id ~ '^[0-9]{15,22}$'),
  source_sha256 text not null check (source_sha256 ~ '^[a-f0-9]{64}$'),
  source_filename text not null,
  source_url text not null check (source_url ~ '^https://(www[.])?kalodata[.]com/(product|shop)/detail[?]'),
  source_rank integer not null check (source_rank>0),
  payload jsonb not null check (jsonb_typeof(payload)='object'),
  imported_at timestamptz not null default now(),
  primary key(market,source_scope,source_entity_id,report_date_range,source_record_id,source_sha256)
);
comment on table public.kalodata_targeted_report_rows is
  'Private source observations from product/shop detail exports. Scope-local sales/ranks never replace global report metrics. Separate SHA256 versions retain provenance.';
alter table public.kalodata_targeted_report_rows enable row level security;
revoke all on public.kalodata_targeted_report_rows from public,anon,authenticated;
grant all on public.kalodata_targeted_report_rows to service_role;
'''


def sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def apply_sql(exports):
    # Shops first so newly verified products can be joined by later video files.
    exports = sorted(exports, key=lambda x: x['scope'] != 'shop_products')
    data = sql_literal(json.dumps(exports, ensure_ascii=False))
    delimiter = '$targeted_' + hashlib.sha256(data.encode()).hexdigest()[:24] + '$'
    while delimiter in data:
        delimiter = delimiter[:-1] + '_x$'
    template = '''-- REVIEW ONLY. This script is generated; the preparer never executes it.
-- Does not call the complete-ranking import RPC, start downloads, or run AI.
-- Existing commercial metrics, global ranks, media and manual links are preserved.
begin;
set local lock_timeout='10s';
set local standard_conforming_strings=on;
select pg_advisory_xact_lock(20261010);
do $targeted$
declare source jsonb; row_data jsonb; item jsonb; old_video public.videos;
  target_product public.products; existing_id uuid; creator uuid; identities integer;
  name_value text; inserted_videos integer:=0; linked_videos integer:=0;
  inserted_products integer:=0; completed_products integer:=0; archived integer:=0;
  columns_sql text; values_sql text; updates_sql text;
  product_columns text[]:=string_to_array('tiktok_product_id producto_nombre producto_url imagen_url categoria precio_mxn price currency total_ventas total_ingresos_mxn revenue_30d gmv_30d_mxn commission commission_amount creators_count creators_active_30d rating market identity_status',' ');
begin
 for source in select value from jsonb_array_elements(__DATA__::jsonb) loop
  if source->>'scope'='shop_products' and not exists(
    select 1 from public.catalog_shops where market=source->>'market' and shop_id=source->>'source_id'
  ) then raise exception 'Source shop is not in the reviewed catalogue: %',source->>'source_id'; end if;
  for row_data in select value from jsonb_array_elements(source->'rows') loop
   item:=row_data->'normalized';
   insert into public.kalodata_targeted_report_rows(
    market,source_scope,source_entity_id,report_date_range,source_record_id,source_sha256,
    source_filename,source_url,source_rank,payload
   ) values(source->>'market',source->>'scope',source->>'source_id',source->>'period',
    row_data->>'source_record_id',source->>'source_sha256',source->>'source_filename',source->>'source_url',
    (row_data->>'source_rank')::integer,row_data||jsonb_build_object('intro',source->'intro'))
   on conflict do nothing;
   archived:=archived+1;
   if source->>'scope'='shop_products' then
    select count(*),(array_agg(id))[1] into identities,existing_id from public.products
     where tiktok_product_id=item->>'tiktok_product_id';
    if identities>1 then raise exception 'Ambiguous product ID: %',item->>'tiktok_product_id'; end if;
    target_product:=null;
    if existing_id is not null then select * into target_product from public.products where id=existing_id; end if;
    if existing_id is not null and (target_product.market is distinct from source->>'market'
      or target_product.identity_status='legacy_duplicate') then raise exception 'Product market/identity conflict'; end if;
    if existing_id is null or target_product.identity_status='video_source' then
     name_value:=item->>'producto_nombre';
     if exists(select 1 from public.products where producto_nombre=name_value and id is distinct from existing_id) then
      item:=item||jsonb_build_object('producto_nombre',name_value||' [ID:'||(item->>'tiktok_product_id')||']');
     end if;
     select string_agg(format('%I',key),','),string_agg(format('s.%I',key),','),string_agg(format('%I=s.%I',key,key),',')
      into columns_sql,values_sql,updates_sql from jsonb_object_keys(item) key where key=any(product_columns);
     if existing_id is null then
      execute format('insert into public.products(%s,rank,report_date_range,last_import,last_imported_from_kalodata_at) select %s,null,$2,now(),now() from jsonb_populate_record(null::public.products,$1) s returning id',columns_sql,values_sql)
       into existing_id using item,source->>'period';
      inserted_products:=inserted_products+1;
     else
      execute format('update public.products t set %s,report_date_range=$3,last_import=now(),last_imported_from_kalodata_at=now() from jsonb_populate_record(null::public.products,$1) s where t.id=$2',updates_sql)
       using item,existing_id,source->>'period';
      completed_products:=completed_products+1;
     end if;
    end if;
    insert into public.catalog_shop_products(market,shop_id,product_id,source_tiktok_product_id,source_url)
     values(source->>'market',source->>'source_id',existing_id,item->>'tiktok_product_id',source->>'source_url')
     on conflict(market,shop_id,product_id) do nothing;
   else
    select count(*),(array_agg(id))[1] into identities,existing_id from public.products
     where tiktok_product_id=source->>'source_id';
    if identities<>1 then raise exception 'Product must exist once before targeted videos: %',source->>'source_id'; end if;
    select * into target_product from public.products where id=existing_id;
    if target_product.market is distinct from source->>'market' or target_product.identity_status='legacy_duplicate'
      or item->>'tiktok_product_id' is distinct from source->>'source_id' then raise exception 'Target product identity conflict'; end if;
    select count(*),(array_agg(id))[1] into identities,existing_id from public.videos
     where tiktok_video_id=item->>'tiktok_video_id' or video_url=item->>'video_url';
    if identities>1 then raise exception 'Ambiguous video ID/URL: %',item->>'tiktok_video_id'; end if;
    if existing_id is null then
     select case when count(*)=1 then (array_agg(id))[1] end into creator from public.creators
      where country=source->>'market' and lower(ltrim(usuario_creador,'@'))=item->>'creator_handle';
     insert into public.videos(tiktok_video_id,video_url,title,creator_handle,creator_id,product_name,category,country,duration,
      tiktok_product_id,source_product_url,product_id,match_status,match_method,match_source,match_confidence,match_algorithm_version,
      rank,revenue_mxn,sales,views,roas,snapshot_date_range,imported_at)
     values(item->>'tiktok_video_id',item->>'video_url',item->>'title',item->>'creator_handle',creator,
      item->>'product_name',item->>'category',source->>'market',(item->>'duration')::integer,
      source->>'source_id',item->>'source_product_url',target_product.id,'matched','tiktok_product_id','exact',1,'v6-id-only',
      null,null,null,null,null,null,now());
     inserted_videos:=inserted_videos+1;
    else
     select * into old_video from public.videos where id=existing_id;
     if (old_video.tiktok_video_id is not null and old_video.tiktok_video_id is distinct from item->>'tiktok_video_id')
       or (old_video.country is not null and old_video.country is distinct from source->>'market') then
      raise exception 'Existing video identity/market conflict: %',item->>'tiktok_video_id';
     end if;
     -- Legacy rows may already contain this exact canonical URL/media without
     -- a separate TikTok ID. Recover the verified ID instead of redownloading.
     if old_video.tiktok_video_id is null then
      if old_video.video_url is distinct from item->>'video_url' then
       raise exception 'Legacy video has no exact canonical URL match';
      end if;
      update public.videos set tiktok_video_id=item->>'tiktok_video_id' where id=existing_id;
     end if;
     -- Other source SKUs/manual links stay untouched. Their scoped relationship
     -- remains in the private archive for an explicit multi-product model later.
     if old_video.manual_match is not true and old_video.product_id is null
       and (old_video.tiktok_product_id is null or old_video.tiktok_product_id=source->>'source_id') then
      update public.videos set product_id=target_product.id,tiktok_product_id=source->>'source_id',
       source_product_url=coalesce(source_product_url,item->>'source_product_url'),country=coalesce(country,source->>'market'),
       match_status='matched',match_method='tiktok_product_id',match_source='exact',match_confidence=1,match_algorithm_version='v6-id-only'
       where id=existing_id;
      linked_videos:=linked_videos+1;
     end if;
    end if;
   end if;
  end loop;
 end loop;
 raise notice 'Scoped observations %, new videos %, exact links filled %, new products %, source-only products completed %',
  archived,inserted_videos,linked_videos,inserted_products,completed_products;
end $targeted$;
commit;
'''
    return template.replace('$targeted$', delimiter).replace('__DATA__', data)


def preflight_sql(exports):
    identities = {row['source_record_id']: row['normalized']['video_url']
                  for source in exports if source['scope']=='product_videos' for row in source['rows']}
    value = sql_literal(json.dumps([dict(tiktok_video_id=key,canonical_url=url)
                                   for key,url in sorted(identities.items())]))
    return f'''-- Read-only exact-ID preflight. Do this again immediately before paid downloads.
with requested as (select value->>'tiktok_video_id' as tiktok_video_id,value->>'canonical_url' as canonical_url
 from jsonb_array_elements({value}::jsonb))
select r.tiktok_video_id,v.id,v.country,v.product_id,v.tiktok_product_id,v.manual_match,v.rank,
 v.revenue_mxn,v.sales,v.views,v.roas,v.snapshot_date_range,v.video_mp4_url,v.analysis_ready,
 d.status as download_status
from requested r left join public.videos v on v.tiktok_video_id=r.tiktok_video_id or v.video_url=r.canonical_url
left join public.catalog_download_jobs d on d.tiktok_video_id=r.tiktok_video_id
order by r.tiktok_video_id;
'''


def application_batches(exports, batch_size=300, max_bytes=900_000):
    """Bound statements by both observation count and actual UTF-8 SQL size."""
    batches = []
    current = []
    current_count = 0
    for source in sorted(exports, key=lambda x: x['scope'] != 'shop_products'):
        for row in source['rows']:
            metadata = {k:v for k,v in source.items() if k!='rows'}
            candidate = [{**part, 'rows':list(part['rows'])} for part in current]
            if candidate and all(candidate[-1][key]==metadata[key] for key in ('scope','source_id','source_sha256')):
                candidate[-1]['rows'].append(row)
            else:
                candidate.append({**metadata, 'rows':[row]})
            if current and (current_count>=batch_size or len(apply_sql(candidate).encode())>max_bytes):
                batches.append(current)
                candidate=[{**metadata,'rows':[row]}]
                current_count=0
            if len(apply_sql(candidate).encode())>max_bytes:
                raise ValueError('One observation exceeds the SQL byte limit; review this source row manually')
            current=candidate
            current_count+=1
    if current:
        batches.append(current)
    return batches


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--manifest', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--market', choices=['mx','us'], default='mx')
    parser.add_argument('--node', default='node')
    parser.add_argument('--batch-size',type=int,default=300)
    parser.add_argument('--max-sql-bytes',type=int,default=900_000)
    args = parser.parse_args()
    if not 1<=args.batch_size<=300 or not 50_000<=args.max_sql_bytes<1_000_000:
        raise ValueError('Batch size must be 1..300 and SQL byte limit 50,000..999,999')
    manifest_path = args.manifest.resolve(strict=True)
    entries = json.loads(manifest_path.read_text())
    if not isinstance(entries, list) or not entries:
        raise ValueError('Manifest must be a non-empty JSON array')
    exports = [prepare_export(entry,manifest_path.parent,args.market,args.node) for entry in entries]
    identities = [(e['scope'],e['source_id'],e['source_sha256']) for e in exports]
    if len(identities) != len(set(identities)):
        raise ValueError('The same scoped file appears twice in the manifest')
    args.out.mkdir(parents=True, exist_ok=True)
    dump(args.out/'normalized.json', exports)
    dump(args.out/'manifest.json', [{k:v for k,v in e.items() if k!='rows'} | {'rows':len(e['rows'])} for e in exports])
    (args.out/'migration.sql').write_text(MIGRATION)
    (args.out/'apply.sql').write_text(apply_sql(exports))
    (args.out/'preflight.sql').write_text(preflight_sql(exports))
    batch_dir=args.out/'apply-batches'
    batch_dir.mkdir(exist_ok=True)
    batches=application_batches(exports,args.batch_size,args.max_sql_bytes)
    batch_manifest=[]
    for index,batch in enumerate(batches,start=1):
        statement=apply_sql(batch)
        filename=f'apply-{index:03d}.sql'
        (batch_dir/filename).write_text(statement)
        batch_manifest.append(dict(filename=filename,rows=sum(len(e['rows']) for e in batch),
                                   bytes=len(statement.encode()),sha256=hashlib.sha256(statement.encode()).hexdigest()))
    expected={x['filename'] for x in batch_manifest}
    for stale in batch_dir.glob('apply-*.sql'):
        if re.fullmatch(r'apply-\d{3,}[.]sql',stale.name) and stale.name not in expected:
            stale.unlink()
    dump(batch_dir/'manifest.json',batch_manifest)
    # Candidates are intentionally NOT an Apify input: existing media and failed
    # permanent downloads need a live DB check before any billable actor starts.
    candidates = {}
    for e in exports:
        if e['scope'] != 'product_videos':
            continue
        for row in e['rows']:
            n=row['normalized'];vid=n['tiktok_video_id']
            c=candidates.setdefault(vid,dict(video_id=vid,url=n['video_url'],source_product_ids=[],requires_live_db_check=True))
            if e['source_id'] not in c['source_product_ids']:
                c['source_product_ids'].append(e['source_id'])
    dump(args.out/'download-candidates-unchecked.json',list(candidates.values()))
    summary=dict(files=len(exports),source_rows=sum(len(e['rows']) for e in exports),
                 unique_video_ids=len(candidates),scopes=sorted({e['scope'] for e in exports}),
                 periods=sorted({e['period'] for e in exports}),
                 global_rank_updates=0,existing_video_metric_updates=0,
                 application_batches=len(batches),largest_sql_bytes=max(x['bytes'] for x in batch_manifest),
                 paid_downloads_started=0,database_writes=0)
    dump(args.out/'summary.json',summary)
    print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
